"""Generate fictional KlassApp demo schools (pre-onboarding checklists)."""
import csv, os, random, sys
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

OUT = sys.argv[1] if len(sys.argv) > 1 else '/workspace/demo-gen/out'

FIRST_F = ["Amina","Grace","Sarah","Esther","Joy","Martha","Ruth","Faith","Mercy","Doreen","Brenda","Agnes","Harriet",
           "Patience","Prossy","Juliet","Winnie","Irene","Sharon","Rebecca","Lydia","Hope","Florence","Annet","Rose",
           "Phiona","Shamim","Asha","Gloria","Vivian","Edith","Peace","Rachel","Naomi","Damalie","Allen","Betty","Carol",
           "Diana","Eva","Hellen","Janet","Lilian","Monica","Olivia","Pauline","Rita","Stella","Teddy","Zaituni"]
FIRST_M = ["Brian","Daniel","Francis","Henry","Peter","Robert","David","Paul","John","Joseph","Moses","Isaac","Samuel",
           "Emmanuel","Ivan","Allan","Denis","Ronald","Solomon","Timothy","Andrew","Benon","Collins","Derrick","Edgar",
           "Fred","Gerald","Hassan","Ibrahim","Jonah","Kenneth","Lawrence","Martin","Nathan","Oscar","Patrick","Richard",
           "Simon","Tonny","Vincent","Wilson","Yusuf","Arthur","Bosco","Charles","Elijah","George","Herbert","Jacob","Mark"]
LAST = ["Nakato","Okello","Nampiima","Wasswa","Kyomuhangi","Ssekitoleko","Mbabazi","Mugisha","Atuhaire","Tumusiime",
        "Nabirye","Kato","Namukasa","Ssali","Achieng","Opio","Akello","Owor","Nalwoga","Byaruhanga","Tugume","Kemigisha",
        "Asiimwe","Ainembabazi","Turyahabwe","Nankya","Lubega","Mukasa","Nakimuli","Kiggundu","Ochieng","Odongo","Apio",
        "Ayebare","Kansiime","Musoke","Nanyonga","Sserwadda","Kigozi","Namutebi","Bwambale","Masereka","Kabugo","Ndyanabo",
        "Rukundo","Natukunda","Ahimbisibwe","Ojok","Adong","Lukwago","Wanyama","Wekesa","Nangobi","Isabirye","Mutebi",
        "Nsubuga","Kizza","Babirye","Namusoke","Kasozi"]

TERMS = [("Term 1","2027-02-03","2027-05-02"),("Term 2","2027-05-24","2027-08-22"),("Term 3","2027-09-13","2027-12-05")]

used_names = set()
used_emails = set()

def person(rng, gender=None):
    while True:
        g = gender or rng.choice(["female","male"])
        first = rng.choice(FIRST_F if g == "female" else FIRST_M)
        last = rng.choice(LAST)
        full = f"{first} {last}"
        if full not in used_names:
            used_names.add(full)
            return full, g

def email_for(name):
    base = name.lower().replace(" ", ".").replace("'", "")
    e = f"{base}@example.com"; n = 2
    while e in used_emails:
        e = f"{base}{n}@example.com"; n += 1
    used_emails.add(e)
    return e

def phone(n):
    # Obviously fake: +256 700 000 xxx (passes the Uganda 9-digit national rule).
    assert 0 <= n <= 999
    return f"+256700000{n:03d}"

def dob(rng, year):
    d = date(year, 1, 1) + timedelta(days=rng.randrange(0, 365))
    return d.isoformat()

CANDIDATE = {"Primary Seven","Senior Four","Senior Six"}

def build_school(cfg):
    rng = random.Random(cfg["seed"])
    school = dict(cfg)
    # Classes and streams
    cls_rows = []
    for cls, streams in cfg["classes"]:
        if streams:
            for s in streams:
                cls_rows.append({"Class": cls, "Stream": s})
        else:
            cls_rows.append({"Class": cls, "Stream": ""})
    school["classes_rows"] = cls_rows
    # Subjects
    school["subject_rows"] = [{"Class": c, "Seeded by KlassApp": ", ".join(seeded), "Add in the Wizard": ", ".join(add)}
                              for c, seeded, add in cfg["subjects"]]
    # Teachers
    t_rows = []
    pn = cfg["teacher_phone_start"]
    for i, (subjects, classes) in enumerate(cfg["teachers"]):
        name, _ = person(rng)
        t_rows.append({"Name": name, "Email": email_for(name), "Subjects": ", ".join(subjects),
                       "Classes": ", ".join(classes), "Phone": phone(pn + i)})
    school["teacher_rows"] = t_rows
    # Class teachers (reuse teachers from the teacher file by index)
    ct_rows = []
    for (cls, stream, tidx) in cfg["class_teachers"]:
        t = t_rows[tidx]
        ct_rows.append({"class": cls, "stream": stream, "name": t["Name"], "email": t["Email"], "phone": t["Phone"]})
    school["class_teacher_rows"] = ct_rows
    # Students
    s_rows = []
    pp = cfg["parent_phone_start"]
    k = 0
    for (cls, stream, count, birth_year) in cfg["students"]:
        for j in range(count):
            name, g = person(rng)
            surname = name.split(" ")[1]
            pg = rng.choice(["female","male"])
            pfirst = rng.choice(FIRST_F if pg == "female" else FIRST_M)
            parent = f"{pfirst} {surname}"
            k += 1
            s_rows.append({
                "Name": name, "Class": cls, "Stream": stream, "Gender": g,
                "Parent Name": parent, "Parent Phone": phone(pp + k - 1),
                "School Student ID": f"{cfg['id_prefix']}-2027-{k:03d}",
                "LIN": "",
                "UNEB Reg No.": (f"U0000/{k:03d}" if cls in CANDIDATE and not stream else ""),
                "Date of Birth": dob(rng, birth_year),
            })
    school["student_rows"] = s_rows
    school["fee_rows"] = cfg["fees"]
    return school

TEACHER_HEADERS = ["Name","Email","Subjects","Classes","Phone"]          # public/templates/teacher-upload-template.xlsx
CLASS_TEACHER_HEADERS = ["class","stream","name","email","phone"]       # public/templates/class-teacher-upload-template.csv
STUDENT_HEADERS = ["Name","Class","Stream","Gender","Parent Name","Parent Phone","School Student ID","LIN","UNEB Reg No.","Date of Birth"]  # StudentUploadTemplateService
CLASSES_HEADERS = ["Class","Stream"]
SUBJECT_HEADERS = ["Class","Seeded by KlassApp","Add in the Wizard"]
FEE_HEADERS = ["Fee name","Amount (UGX)","Applies to","Class","Term"]

TABLES = [  # (sheet, csv file, headers, key)
    ("Classes and Streams", "classes-and-streams.csv", CLASSES_HEADERS, "classes_rows"),
    ("Subjects", "subjects.csv", SUBJECT_HEADERS, "subject_rows"),
    ("Teacher Upload", "teacher-upload-template.csv", TEACHER_HEADERS, "teacher_rows"),
    ("Class Teacher Upload", "class-teacher-upload-template.csv", CLASS_TEACHER_HEADERS, "class_teacher_rows"),
    ("Students", "student-upload-template.csv", STUDENT_HEADERS, "student_rows"),
    ("Fees", "fees.csv", FEE_HEADERS, "fee_rows"),
]

def checklist(s):
    rows = [
        ("School name", "School name", s["name"]),
        ("School size", "Approximate number of students", "Up to 500"),
        ("Country", "Country", "Uganda"),
        ("Curriculum", "Board / Curriculum", "UNEB (Uganda National Examinations Board)"),
        ("School category", "School category", s["category_label"]),
        ("EMIS", "EMIS / Ministry code", s["emis"]),
        ("UNEB centre", "UNEB centre number", "Leave blank (skip)"),
        ("Academic year", "Description", "2027"),
        ("Academic year", "Start date", "2027-01-01"),
        ("Academic year", "End date", "2027-12-31"),
        ("Classes and streams", "Classes", s["classes_summary"]),
        ("Subjects", "Subjects", "See the Subjects sheet"),
        ("Teachers", "Teacher file", f"teacher-upload-template.csv ({len(s['teacher_rows'])} teachers)"),
        ("Teachers", "Class teacher file", f"class-teacher-upload-template.csv ({len(s['class_teacher_rows'])} class teachers)"),
        ("Students", "Student file", f"student-upload-template.csv ({len(s['student_rows'])} students)"),
    ]
    for (n, a, b) in TERMS:
        rows.append(("Terms", n, f"{a} to {b}"))
    rows += [
        ("Terms", "Current term", "Term 1"),
        ("Fees", "Fees", f"See the Fees sheet ({len(s['fee_rows'])} fees)"),
        ("Admin", "Admin name", s["admin"][0]),
        ("Admin", "Admin email", s["admin"][1]),
        ("Admin", "Admin phone", s["admin"][2]),
        ("WhatsApp", "School WhatsApp number", s["whatsapp"]),
        ("Plan", "Plan", "Comes from the invite (Growth, free for 3 months); no choice needed"),
    ]
    return rows

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="0F766E")

def write_sheet(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = HDR_FONT; c.fill = HDR_FILL
    for r in rows:
        ws.append([str(r.get(h, "")) for h in headers])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.number_format = "@"  # keep phones, dates and IDs as text
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r.get(h, ""))) for r in rows]) + 2
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(width, 10), 70)
    ws.freeze_panes = "A2"

def write_school(s):
    d = os.path.join(OUT, s["slug"]); os.makedirs(d, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Checklist"
    ws.append(["Wizard step", "Field", "Value"])
    for c in ws[1]:
        c.font = HDR_FONT; c.fill = HDR_FILL
    for r in checklist(s):
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.number_format = "@"; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 32; ws.column_dimensions["C"].width = 80
    ws.freeze_panes = "A2"
    for sheet, fname, headers, key in TABLES:
        write_sheet(wb.create_sheet(sheet), headers, s[key])
        with open(os.path.join(d, fname), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f, lineterminator="\n")
            w.writerow(headers)
            for r in s[key]:
                w.writerow([r.get(h, "") for h in headers])
    wb.active = 0
    wb.save(os.path.join(d, f"{s['slug']}.xlsx"))

# ---------------- Sunrise Primary and Nursery School ----------------
P = ["Primary One","Primary Two","Primary Three","Primary Four","Primary Five","Primary Six","Primary Seven"]
N = ["Baby Class","Middle Class","Top Class"]
NURS_SEED = ["Literacy","Numeracy","Motor Skills","Social/Emotional"]
PRI_SEED = ["English Language","Mathematics","Social Studies","Science"]
NURS_ALL = NURS_SEED + ["Creative Arts"]

def fee(name, amount, cls, term):
    return {"Fee name": name, "Amount (UGX)": str(amount), "Applies to": "One class" if cls else "Whole school",
            "Class": cls, "Term": term}

sun_tuition = {"Baby Class":250000,"Middle Class":250000,"Top Class":280000,"Primary One":320000,"Primary Two":320000,
               "Primary Three":340000,"Primary Four":360000,"Primary Five":380000,"Primary Six":400000,"Primary Seven":450000}
sun_fees = [fee(f"Tuition {t}", sun_tuition[c], c, t) for c in N + P for (t, _, _) in TERMS]
sun_fees += [fee(f"Lunch {t}", 120000, "", t) for (t, _, _) in TERMS]
sun_fees += [fee("Development Fund", 50000, "", "Yearly"), fee("PLE Registration", 45000, "Primary Seven", "Yearly")]

SUNRISE = dict(
    seed=2027_1, slug="sunrise-primary-and-nursery-school", name="Sunrise Primary and Nursery School",
    category_label="Primary + Nursery", category="primary_nursery", emis="DEMO-0001", id_prefix="SUN",
    admin=("Grace Atuhaire", "admin.sunrise@example.com", phone(300)), whatsapp=phone(301),
    classes=[("Baby Class",[]),("Middle Class",[]),("Top Class",[]),("Primary One",["A","B"]),("Primary Two",[]),
             ("Primary Three",["A","B"]),("Primary Four",[]),("Primary Five",["A","B"]),("Primary Six",[]),("Primary Seven",[])],
    classes_summary="Baby, Middle and Top Class; Primary One to Seven. Streams A and B on Primary One, Three and Five.",
    subjects=[(c, NURS_SEED, ["Creative Arts"]) for c in N]
             + [(c, PRI_SEED, ["Religious Education","Kiswahili","Physical Education"]) for c in P[:3]]
             + [(c, PRI_SEED, ["Religious Education","Kiswahili"]) for c in P[3:]],
    teacher_phone_start=310,
    teachers=[(NURS_ALL,["Baby Class"]),(NURS_ALL,["Middle Class"]),(NURS_ALL,["Top Class"]),
              (["English Language"],P[0:2]),(["Mathematics"],P[0:2]),(["English Language"],P[2:4]),(["Mathematics"],P[2:4]),
              (["English Language"],P[4:7]),(["Mathematics"],P[4:7]),(["Science"],P[0:4]),(["Science"],P[4:7]),
              (["Social Studies"],P),(["Religious Education"],P),(["Kiswahili"],P),(["Physical Education"],P[0:3])],
    class_teachers=[("Baby Class","",0),("Middle Class","",1),("Top Class","",2),("Primary One","A",3),("Primary One","B",4),
                    ("Primary Two","",9),("Primary Three","A",5),("Primary Three","B",6),("Primary Four","",11),
                    ("Primary Five","A",7),("Primary Five","B",8),("Primary Six","",10),("Primary Seven","",12)],
    parent_phone_start=400,
    students=[("Baby Class","",10,2023),("Middle Class","",10,2022),("Top Class","",10,2021),
              ("Primary One","A",9,2020),("Primary One","B",9,2020),("Primary Two","",10,2019),
              ("Primary Three","A",8,2018),("Primary Three","B",8,2018),("Primary Four","",10,2017),
              ("Primary Five","A",8,2016),("Primary Five","B",8,2016),("Primary Six","",10,2015),("Primary Seven","",10,2014)],
    fees=sun_fees,
)

# ---------------- Hilltop O and A Level School ----------------
O = ["Senior One","Senior Two","Senior Three","Senior Four"]
A = ["Senior Five","Senior Six"]
O_SEED = ["English Language","General Mathematics","Biology","Chemistry","Physics","Geography","History and Political Education"]
O_ADD = ["Christian Religious Education","Agriculture","Kiswahili","Physical Education","Entrepreneurship Education","ICT"]
A_SEED = ["General Paper","Biology","Chemistry","Physics","History and Political Education","Literature in English","ICT",
          "General Mathematics","Geography","Economics","Entrepreneurship Education"]
A_ADD = ["Mathematics"]
COMBOS = ["PCM","BCM","HEG","MEG"]

hil_tuition = {"Senior One":650000,"Senior Two":650000,"Senior Three":700000,"Senior Four":750000,"Senior Five":800000,"Senior Six":850000}
hil_fees = [fee(f"Tuition {t}", hil_tuition[c], c, t) for c in O + A for (t, _, _) in TERMS]
hil_fees += [fee(f"Laboratory {t}", 60000, "", t) for (t, _, _) in TERMS]
hil_fees += [fee("Development Fund", 100000, "", "Yearly"), fee("UCE Registration", 165000, "Senior Four", "Yearly"),
             fee("UACE Registration", 190000, "Senior Six", "Yearly")]

HILLTOP = dict(
    seed=2027_2, slug="hilltop-o-and-a-level-school", name="Hilltop O and A Level School",
    category_label="O-Level + A-Level", category="o_a_level", emis="DEMO-0002", id_prefix="HIL",
    admin=("Samuel Byaruhanga", "admin.hilltop@example.com", phone(600)), whatsapp=phone(601),
    classes=[(c, ["A","B"]) for c in O] + [(c, COMBOS) for c in A],
    classes_summary="Senior One to Four with streams A and B; Senior Five and Six with A level combinations PCM, BCM, HEG and MEG as streams.",
    subjects=[(c, O_SEED, O_ADD) for c in O] + [(c, A_SEED, A_ADD) for c in A],
    teacher_phone_start=610,
    teachers=[(["English Language"],O[0:2]),(["English Language"],O[2:4]),(["General Mathematics"],O[0:2]),
              (["General Mathematics"],O[2:4]),(["Biology"],O),(["Chemistry"],O),(["Physics"],O),(["Geography"],O),
              (["History and Political Education"],O),(["Christian Religious Education"],O),(["Agriculture"],O),
              (["Kiswahili"],O),(["Physical Education"],O),(["Entrepreneurship Education"],O),(["ICT"],O),
              (["General Paper"],A),(["Biology"],A),(["Chemistry"],A),(["Physics"],A),(["Mathematics"],A),(["Economics"],A),
              (["Geography"],A),(["History and Political Education"],A),(["Literature in English"],A),
              (["ICT","General Mathematics","Entrepreneurship Education"],A)],
    class_teachers=[(c, s, i*2 + j) for i, c in enumerate(O) for j, s in enumerate(["A","B"])]
                   + [(c, s, 16 + i*4 + j if 16 + i*4 + j < 25 else 15) for i, c in enumerate(A) for j, s in enumerate(COMBOS)],
    parent_phone_start=700,
    students=[("Senior One","A",20,2014),("Senior One","B",20,2014),("Senior Two","A",18,2013),("Senior Two","B",18,2013),
              ("Senior Three","A",17,2012),("Senior Three","B",17,2012),("Senior Four","A",15,2011),("Senior Four","B",15,2011),
              ("Senior Five","PCM",8,2010),("Senior Five","BCM",7,2010),("Senior Five","HEG",8,2010),("Senior Five","MEG",7,2010),
              ("Senior Six","PCM",8,2009),("Senior Six","BCM",7,2009),("Senior Six","HEG",7,2009),("Senior Six","MEG",8,2009)],
    fees=hil_fees,
)

if __name__ == "__main__":
    for cfg in (SUNRISE, HILLTOP):
        s = build_school(cfg)
        write_school(s)
        print(s["slug"], {k: len(s[k]) for _, _, _, k in TABLES})
